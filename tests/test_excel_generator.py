import pytest
import io
from openpyxl import load_workbook
from unittest.mock import MagicMock
from apps.parser.excel_generator import ExcelGenerator


def make_mock_uc(order, uc_id, description, is_automated=False, status='À tester'):
    uc = MagicMock()
    uc.order = order
    uc.use_case_text = uc_id
    uc.description = description
    uc.preconditions = 'Prérequis'
    uc.steps = 'Étapes de test'
    uc.expected_results = 'Résultat attendu'
    uc.observed_results = ''
    uc.is_automated = is_automated
    uc.status = status
    uc.jira_ticket = ''
    return uc


def get_data_rows(ws, min_row=3):
    """Return rows with numeric first cell (actual UC data rows, not summary)."""
    return [
        row for row in ws.iter_rows(min_row=min_row)
        if row[0].value and isinstance(row[0].value, int)
    ]


class TestExcelGenerator:
    def test_generates_two_sheets(self):
        ucs = [make_mock_uc(1, 'UC001', 'Test login')]
        gen = ExcelGenerator(ucs, 'test.docx')
        buf = gen.generate()
        wb = load_workbook(buf)
        assert 'Use Cases' in wb.sheetnames
        assert 'Cas automatisé' in wb.sheetnames

    def test_use_cases_sheet_has_correct_count(self):
        ucs = [make_mock_uc(i, f'UC{i:03d}', f'Description {i}') for i in range(1, 6)]
        gen = ExcelGenerator(ucs, 'test.docx')
        buf = gen.generate()
        wb = load_workbook(buf)
        ws = wb['Use Cases']
        data_rows = get_data_rows(ws)
        assert len(data_rows) == 5

    def test_automated_sheet_filters_correctly(self):
        ucs = [
            make_mock_uc(1, 'UC001', 'Manuel', is_automated=False),
            make_mock_uc(2, 'UC002', 'Auto 1', is_automated=True),
            make_mock_uc(3, 'UC003', 'Auto 2', is_automated=True),
        ]
        gen = ExcelGenerator(ucs, 'test.docx')
        buf = gen.generate()
        wb = load_workbook(buf)
        ws = wb['Cas automatisé']
        data_rows = get_data_rows(ws)
        assert len(data_rows) == 2

    def test_empty_use_cases(self):
        gen = ExcelGenerator([], 'empty.docx')
        buf = gen.generate()
        wb = load_workbook(buf)
        assert 'Use Cases' in wb.sheetnames
        assert 'Cas automatisé' in wb.sheetnames

    def test_returns_bytes_io(self):
        ucs = [make_mock_uc(1, 'UC001', 'Test')]
        gen = ExcelGenerator(ucs, 'test.docx')
        buf = gen.generate()
        assert isinstance(buf, io.BytesIO)
        assert buf.tell() == 0  # seeked to start

    def test_header_row_content(self):
        ucs = [make_mock_uc(1, 'UC001', 'Test')]
        gen = ExcelGenerator(ucs, 'test.docx')
        buf = gen.generate()
        wb = load_workbook(buf)
        ws = wb['Use Cases']
        headers = [cell.value for cell in ws[3]]
        assert headers[0] == 'N°'
        assert 'CAS' in headers
        assert 'Use Case' in headers
        assert 'Description' in headers
        assert 'Résultats Attendus' in headers

    def test_cas_column_format(self):
        ucs = [
            make_mock_uc(1, 'Login test', 'Test de connexion'),
            make_mock_uc(42, 'Logout test', 'Test de déconnexion'),
        ]
        gen = ExcelGenerator(ucs, 'test.docx')
        buf = gen.generate()
        wb = load_workbook(buf)
        ws = wb['Use Cases']
        rows = [row for row in ws.iter_rows(min_row=4, max_row=5) if row[0].value]
        assert rows[0][0].value == 1      # N° (row number)
        assert rows[0][1].value == 'UC-001'  # CAS from uc.order
        assert rows[0][2].value in (None, '')  # Tickets Jira (vide par défaut)
        assert rows[0][3].value == 'Login test'  # Use Case
        assert rows[1][0].value == 2      # N° (row number = index)
        assert rows[1][1].value == 'UC-042'  # CAS from uc.order
        assert rows[1][3].value == 'Logout test'

    def test_cas_in_automated_sheet(self):
        ucs = [
            make_mock_uc(1, 'Auto test', 'Automated', is_automated=True),
            make_mock_uc(2, 'Manual test', 'Manual', is_automated=False),
        ]
        gen = ExcelGenerator(ucs, 'test.docx')
        buf = gen.generate()
        wb = load_workbook(buf)
        ws = wb['Cas automatisé']
        rows = [row for row in ws.iter_rows(min_row=4) if isinstance(row[0].value, int)]
        assert len(rows) == 1
        assert rows[0][1].value == 'UC-001'
        assert rows[0][3].value == 'Auto test'

    def test_source_filename_in_title(self):
        ucs = [make_mock_uc(1, 'UC001', 'Test')]
        gen = ExcelGenerator(ucs, 'ma_recette.docx')
        buf = gen.generate()
        wb = load_workbook(buf)
        ws = wb['Use Cases']
        meta_cell = ws.cell(row=2, column=1).value
        assert 'ma_recette.docx' in (meta_cell or '')

    def test_summary_row_exists(self):
        ucs = [make_mock_uc(i, f'UC{i:03d}', f'Desc') for i in range(1, 4)]
        gen = ExcelGenerator(ucs, 'test.docx')
        buf = gen.generate()
        wb = load_workbook(buf)
        ws = wb['Use Cases']
        # Find summary row (contains "Total")
        found = any(
            'Total' in str(cell.value)
            for row in ws.iter_rows()
            for cell in row
            if cell.value
        )
        assert found

    def test_no_automated_in_automated_sheet(self):
        ucs = [make_mock_uc(1, 'UC001', 'Manuel', is_automated=False)]
        gen = ExcelGenerator(ucs, 'test.docx')
        buf = gen.generate()
        wb = load_workbook(buf)
        ws = wb['Cas automatisé']
        data_rows = get_data_rows(ws)
        assert len(data_rows) == 0

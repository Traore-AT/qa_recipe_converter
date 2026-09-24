"""
Tests unitaires — Tickets Jira

Couvre :
- Colonne 'Tickets Jira' après 'CAS' dans la génération Excel + hyperlien
- Exposure dans le board Kanban (jira_ticket + jira_url)
- Endpoint PATCH use-case jira (UseCaseJiraUpdateView)
"""
import uuid
import pytest
from django.test import override_settings
from django.urls import reverse
from rest_framework import status

from apps.core.models import ConversionJob, ExtractedUseCase
from apps.qamanagement.models import UseCaseAssignment
from tests.test_qamanagement import (  # reuse fixtures
    owner, team, project, lead_pm, tester_pm, done_job, sprint, assignment,
    lead_user, tester_user, viewer_user,
)
from tests.test_qamanagement import client


# ─────────────────────────────────────────────────────────────────────────────
# EXCEL
# ─────────────────────────────────────────────────────────────────────────────
class TestExcelJiraColumn:
    def test_column_after_cas_with_hyperlink(self, done_job):
        from apps.parser.excel_generator import ExcelGenerator, COLUMNS

        assert [c[1] for c in COLUMNS[:2]] == ['CAS', 'Tickets Jira']

        job, ucs = done_job
        uc = ucs[0]
        uc.jira_ticket = 'QA-123'
        uc.save()

        with override_settings(JIRA_BASE_URL='https://acme.atlassian.net'):
            gen = ExcelGenerator([uc], source_filename=job.source_filename)
            buf = gen.generate()

        from openpyxl import load_workbook
        wb = load_workbook(buf)
        ws = wb['Use Cases']

        headers = [ws.cell(row=3, column=c).value for c in range(1, 11)]
        assert 'Tickets Jira' in headers
        assert headers.index('Tickets Jira') == headers.index('CAS') + 1

        cell = ws.cell(row=4, column=3)
        assert cell.value == 'QA-123'
        assert cell.hyperlink is not None
        assert cell.hyperlink.target == 'https://acme.atlassian.net/browse/QA-123'

    def test_no_hyperlink_without_config(self, done_job):
        from apps.parser.excel_generator import ExcelGenerator

        job, ucs = done_job
        uc = ucs[0]
        uc.jira_ticket = 'QA-456'
        uc.save()

        with override_settings(JIRA_BASE_URL=''):
            gen = ExcelGenerator([uc], source_filename=job.source_filename)
            buf = gen.generate()

        from openpyxl import load_workbook
        ws = load_workbook(buf)['Use Cases']
        cell = ws.cell(row=4, column=3)
        assert cell.value == 'QA-456'
        assert cell.hyperlink is None


# ─────────────────────────────────────────────────────────────────────────────
# API / SERIALIZER
# ─────────────────────────────────────────────────────────────────────────────
class TestJiraAPI:
    def test_board_exposes_jira_fields(self, client, team, project, sprint, assignment, owner):
        assignment.use_case.jira_ticket = 'QA-111'
        assignment.use_case.save()
        client.force_authenticate(owner)
        with override_settings(JIRA_BASE_URL='https://acme.atlassian.net'):
            url = reverse('qamanagement:sprint-board', args=[team.slug, project.slug, sprint.id])
            resp = client.get(url)
        assert resp.status_code == status.HTTP_200_OK
        col = resp.data['columns']['À tester']
        assert col, 'assignment attendu dans la colonne "À tester"'
        row = col[0]
        assert row['jira_ticket'] == 'QA-111'
        assert row['jira_url'] == 'https://acme.atlassian.net/browse/QA-111'

    def test_update_jira_ticket(self, client, team, project, done_job, owner):
        uc = done_job[1][0]
        client.force_authenticate(owner)
        url = reverse('qamanagement:uc-jira-update', args=[team.slug, project.slug, uc.id])
        resp = client.patch(url, {'jira_ticket': 'QA-999'}, format='json')
        assert resp.status_code == status.HTTP_200_OK
        uc.refresh_from_db()
        assert uc.jira_ticket == 'QA-999'

    def test_update_jira_requires_auth(self, client, team, project, done_job):
        uc = done_job[1][0]
        url = reverse('qamanagement:uc-jira-update', args=[team.slug, project.slug, uc.id])
        resp = client.patch(url, {'jira_ticket': 'QA-1'}, format='json')
        assert resp.status_code in (status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN)
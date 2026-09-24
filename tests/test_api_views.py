"""
Tests unitaires — API Conversion (apps.api.views)

Couvre :
- HealthCheckView, CsrfTokenView
- JobDetailAPIView, JobListAPIView (pagination)
- UseCaseUpdateAPIView (PATCH individuel)
- BatchSaveAPIView (sauvegarde batch des UC + en-tête)
- OpenVSCodeAPIView (génération projet Cypress)
- GenerateExcelAPIView (avec company_name, excel_filename, logo)
- GenerateGherkinAPIView (mode=gherkin et mode=cypress)
- JobDeleteAPIView (suppression recette avec permissions)
"""
import io
import sys
import zipfile
import uuid
from datetime import timedelta

import pytest
from django.utils import timezone
from django.urls import reverse
from django.contrib.auth.models import User
from rest_framework.test import APIClient

from apps.core.models import ConversionJob, ExtractedUseCase
from apps.teams.models import Team, TeamMember, Project, ProjectMember


# ─────────────────────────────────────────────────────────────────────────────
# FIXTURES
# ─────────────────────────────────────────────────────────────────────────────
@pytest.fixture
def api_client():
    return APIClient()


@pytest.fixture
def pending_job(db):
    return ConversionJob.objects.create(
        source_filename='test.docx',
        word_file='uploads/word/test.docx',
        status=ConversionJob.Status.PENDING,
    )


@pytest.fixture
def done_job(db):
    return ConversionJob.objects.create(
        source_filename='recette_v2.docx',
        word_file='uploads/word/recette_v2.docx',
        status=ConversionJob.Status.DONE,
        use_cases_count=3,
        company_name='ACME Corp',
        excel_filename='mon_fichier.xlsx',
    )


@pytest.fixture
def three_use_cases(done_job):
    ucs = []
    for i in range(3):
        uc = ExtractedUseCase.objects.create(
            job=done_job,
            order=i + 1,
            use_case_text=f'UC{i+1:03d}',
            description=f'Cas test {i+1}',
            preconditions='Utilisateur connecté' if i == 0 else '',
            steps=f'Étape {i+1}.1\nÉtape {i+1}.2',
            expected_results=f'Résultat attendu {i+1}',
            is_automated=(i == 1),  # UC002 automatisé
        )
        ucs.append(uc)
    return ucs


@pytest.fixture
def error_job(db):
    return ConversionJob.objects.create(
        source_filename='broken.docx',
        word_file='uploads/word/broken.docx',
        status=ConversionJob.Status.ERROR,
        error_message='Fichier invalide',
    )


# ─────────────────────────────────────────────────────────────────────────────
# HEALTH / CSRF
# ─────────────────────────────────────────────────────────────────────────────
class TestHealthAndCSRF:
    def test_healthcheck_returns_healthy(self, api_client):
        resp = api_client.get('/api/health/')
        assert resp.status_code == 200
        assert resp.data['status'] == 'healthy'

    def test_csrf_sets_cookie(self, api_client):
        resp = api_client.get('/api/csrf/')
        assert resp.status_code == 200
        assert resp.data['csrfToken']
        assert 'csrftoken' in resp.cookies or 'Set-Cookie' in resp


# ─────────────────────────────────────────────────────────────────────────────
# JOB LIST / DETAIL
# ─────────────────────────────────────────────────────────────────────────────
class TestJobListAPI:
    def test_list_returns_paginated_results(self, api_client, pending_job, done_job, error_job):
        resp = api_client.get('/api/jobs/')
        assert resp.status_code == 200
        assert 'results' in resp.data
        assert 'count' in resp.data
        assert resp.data['count'] >= 3
        assert len(resp.data['results']) >= 3

    def test_list_ordered_by_created_at_desc(self, api_client, pending_job):
        old = ConversionJob.objects.create(
            source_filename='old.docx',
            word_file='uploads/word/old.docx',
        )
        # Force old created_at in the past
        old.created_at = timezone.now() - timedelta(hours=2)
        old.save()

        recent = ConversionJob.objects.create(
            source_filename='recent.docx',
            word_file='uploads/word/recent.docx',
        )

        resp = api_client.get('/api/jobs/')
        slugs = [j['source_filename'] for j in resp.data['results']]
        # Most recent first
        assert slugs.index('recent.docx') < slugs.index('old.docx')

    def test_list_pagination_structure(self, api_client, db):
        resp = api_client.get('/api/jobs/?page=1')
        assert resp.status_code == 200
        assert 'count' in resp.data
        assert 'results' in resp.data
        assert 'next' in resp.data or 'previous' in resp.data

    def test_list_empty_returns_empty_results(self, api_client, db):
        ConversionJob.objects.all().delete()
        resp = api_client.get('/api/jobs/')
        assert resp.status_code == 200
        assert resp.data['count'] == 0
        assert resp.data['results'] == []


class TestJobDetailAPI:
    def test_get_job_by_id(self, api_client, done_job):
        resp = api_client.get(f'/api/jobs/{done_job.id}/')
        assert resp.status_code == 200
        assert resp.data['source_filename'] == 'recette_v2.docx'
        assert resp.data['status'] == 'DONE'
        assert resp.data['use_cases_count'] == 3

    def test_get_job_includes_use_cases(self, api_client, done_job, three_use_cases):
        resp = api_client.get(f'/api/jobs/{done_job.id}/')
        assert resp.status_code == 200
        assert len(resp.data['use_cases']) == 3

    def test_get_job_not_found(self, api_client, db):
        resp = api_client.get(f'/api/jobs/{uuid.uuid4()}/')
        assert resp.status_code == 404

    @pytest.mark.skipif(
        sys.version_info >= (3, 14),
        reason='Django 5.0.x incompatible avec Python 3.14 pour la template 404',
    )
    def test_get_job_invalid_uuid(self, api_client, db):
        resp = api_client.get('/api/jobs/not-a-uuid/')
        assert resp.status_code == 404


# ─────────────────────────────────────────────────────────────────────────────
# USE CASE UPDATE
# ─────────────────────────────────────────────────────────────────────────────
class TestUseCaseUpdateAPI:
    def test_patch_status(self, api_client, done_job, three_use_cases):
        uc = three_use_cases[0]
        resp = api_client.patch(
            f'/api/jobs/{done_job.id}/use-cases/{uc.id}/',
            {'status': 'Passé'},
            format='json',
        )
        assert resp.status_code == 200
        assert resp.data['status'] == 'Passé'

    def test_patch_is_automated(self, api_client, done_job, three_use_cases):
        uc = three_use_cases[0]
        assert uc.is_automated is False
        resp = api_client.patch(
            f'/api/jobs/{done_job.id}/use-cases/{uc.id}/',
            {'is_automated': True},
            format='json',
        )
        assert resp.status_code == 200
        assert resp.data['is_automated'] is True

    def test_patch_description(self, api_client, done_job, three_use_cases):
        uc = three_use_cases[0]
        resp = api_client.patch(
            f'/api/jobs/{done_job.id}/use-cases/{uc.id}/',
            {'description': 'Mise à jour test'},
            format='json',
        )
        assert resp.status_code == 200
        assert resp.data['description'] == 'Mise à jour test'

    def test_patch_invalid_uc_id(self, api_client, done_job):
        resp = api_client.patch(
            f'/api/jobs/{done_job.id}/use-cases/{uuid.uuid4()}/',
            {'status': 'Passé'},
            format='json',
        )
        assert resp.status_code == 404

    def test_patch_invalid_status_returns_error(self, api_client, done_job, three_use_cases):
        uc = three_use_cases[0]
        resp = api_client.patch(
            f'/api/jobs/{done_job.id}/use-cases/{uc.id}/',
            {'status': 'INVALID_STATUS'},
            format='json',
        )
        assert resp.status_code == 400


# ─────────────────────────────────────────────────────────────────────────────
# BATCH SAVE
# ─────────────────────────────────────────────────────────────────────────────
class TestBatchSaveAPI:
    def test_batch_save_updates_use_cases(self, api_client, done_job, three_use_cases):
        resp = api_client.post(f'/api/jobs/{done_job.id}/save/', {
            'use_cases': [
                {'id': str(three_use_cases[0].id), 'status': 'Passé', 'is_automated': True},
                {'id': str(three_use_cases[1].id), 'description': 'Descriptions mises à jour'},
                {'id': str(three_use_cases[2].id), 'status': 'Échoué'},
            ],
        }, format='json')
        assert resp.status_code == 200
        assert resp.data['success'] is True
        assert resp.data['automated_count'] == 2  # UC001 now automated + UC002 already automated

        # Verify db persisted
        three_use_cases[0].refresh_from_db()
        assert three_use_cases[0].status == 'Passé'
        assert three_use_cases[0].is_automated is True
        three_use_cases[1].refresh_from_db()
        assert three_use_cases[1].description == 'Descriptions mises à jour'

    def test_batch_save_updates_header_fields(self, api_client, done_job, three_use_cases):
        resp = api_client.post(f'/api/jobs/{done_job.id}/save/', {
            'company_name': 'New Corp',
            'excel_filename': 'export_test.xlsx',
            'use_cases': [],
        }, format='json')
        assert resp.status_code == 200
        done_job.refresh_from_db()
        assert done_job.company_name == 'New Corp'
        assert done_job.excel_filename == 'export_test.xlsx'

    def test_batch_save_empty_use_cases(self, api_client, done_job):
        resp = api_client.post(f'/api/jobs/{done_job.id}/save/', {
            'use_cases': [],
        }, format='json')
        assert resp.status_code == 200
        assert resp.data['success'] is True

    def test_batch_save_invalid_uc_id(self, api_client, done_job):
        resp = api_client.post(f'/api/jobs/{done_job.id}/save/', {
            'use_cases': [{'id': str(uuid.uuid4()), 'status': 'Passé'}],
        }, format='json')
        assert resp.status_code == 404

    def test_batch_save_job_not_found(self, api_client, db):
        resp = api_client.post(f'/api/jobs/{uuid.uuid4()}/save/', {
            'use_cases': [],
        }, format='json')
        assert resp.status_code == 404


# ─────────────────────────────────────────────────────────────────────────────
# VS CODE
# ─────────────────────────────────────────────────────────────────────────────
class TestOpenVSCodeAPI:
    def test_vscode_returns_project_path(self, api_client, done_job, three_use_cases):
        resp = api_client.post(f'/api/jobs/{done_job.id}/vscode/', format='json')
        assert resp.status_code == 200
        assert resp.data['success'] is True
        assert 'project_path' in resp.data
        assert resp.data['automated_count'] == 1  # only UC002 is automated

    def test_vscode_rejects_non_done_job(self, api_client, pending_job):
        resp = api_client.post(f'/api/jobs/{pending_job.id}/vscode/', format='json')
        assert resp.status_code == 404

    def test_vscode_generates_valid_cypress_zip(self, api_client, done_job, three_use_cases):
        resp = api_client.post(f'/api/jobs/{done_job.id}/vscode/', format='json')
        assert resp.status_code == 200
        # Check the generated project contains cypress files
        import os
        project_path = resp.data['project_path']
        assert os.path.isdir(project_path)
        cypress_dir = os.path.join(project_path, 'cypress')
        assert os.path.isdir(cypress_dir), f"Cypress dir missing in {project_path}"

    def test_vscode_job_not_found(self, api_client, db):
        resp = api_client.post(f'/api/jobs/{uuid.uuid4()}/vscode/', format='json')
        assert resp.status_code == 404


# ─────────────────────────────────────────────────────────────────────────────
# GENERATE EXCEL
# ─────────────────────────────────────────────────────────────────────────────
class TestGenerateExcelAPI:
    def test_generate_excel_returns_valid_xlsx(self, api_client, done_job, three_use_cases):
        resp = api_client.get(f'/api/jobs/{done_job.id}/generate/')
        assert resp.status_code == 200
        assert resp['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        assert 'mon_fichier.xlsx' in resp['Content-Disposition']
        assert len(resp.content) > 0
        # Verify it's a valid xlsx (zip-based)
        assert resp.content[:2] == b'PK'

    def test_generate_excel_default_filename(self, api_client, db):
        job = ConversionJob.objects.create(
            source_filename='custom.docx',
            word_file='uploads/word/custom.docx',
            status=ConversionJob.Status.DONE,
            use_cases_count=1,
        )
        ExtractedUseCase.objects.create(job=job, order=1, description='Test')
        resp = api_client.get(f'/api/jobs/{job.id}/generate/')
        assert resp.status_code == 200
        assert 'recette_custom.xlsx' in resp['Content-Disposition']

    def test_generate_excel_rejects_non_done_job(self, api_client, pending_job):
        resp = api_client.get(f'/api/jobs/{pending_job.id}/generate/')
        assert resp.status_code == 404

    def test_generate_excel_job_not_found(self, api_client, db):
        resp = api_client.get(f'/api/jobs/{uuid.uuid4()}/generate/')
        assert resp.status_code == 404


# ─────────────────────────────────────────────────────────────────────────────
# GENERATE GHERKIN
# ─────────────────────────────────────────────────────────────────────────────
class TestGenerateGherkinAPI:
    def test_gherkin_default_mode_returns_zip(self, api_client, done_job, three_use_cases):
        resp = api_client.get(f'/api/jobs/{done_job.id}/gherkin/')
        assert resp.status_code == 200
        assert resp['Content-Type'] == 'application/zip'
        assert 'gherkin_features_' in resp['Content-Disposition']
        assert resp.content[:2] == b'PK'
        # Verify zip contains .feature files
        with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
            names = zf.namelist()
            feature_files = [n for n in names if n.endswith('.feature')]
            assert len(feature_files) > 0

    def test_gherkin_cypress_mode_returns_zip(self, api_client, done_job, three_use_cases):
        resp = api_client.get(f'/api/jobs/{done_job.id}/gherkin/?mode=cypress')
        assert resp.status_code == 200
        assert resp['Content-Type'] == 'application/zip'
        assert 'cypress_project_' in resp['Content-Disposition']
        assert resp.content[:2] == b'PK'
        # Verify zip contains cypress structure
        with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
            names = zf.namelist()
            cypress_configs = [n for n in names if 'cypress.config' in n]
            assert len(cypress_configs) > 0, f"No cypress.config found in {names[:10]}"

    def test_gherkin_invalid_mode_defaults_to_gherkin(self, api_client, done_job, three_use_cases):
        resp = api_client.get(f'/api/jobs/{done_job.id}/gherkin/?mode=invalid')
        assert resp.status_code == 200
        assert resp['Content-Type'] == 'application/zip'
        assert 'gherkin_features_' in resp['Content-Disposition']

    def test_gherkin_rejects_non_done_job(self, api_client, pending_job):
        resp = api_client.get(f'/api/jobs/{pending_job.id}/gherkin/')
        assert resp.status_code == 404

    def test_gherkin_job_not_found(self, api_client, db):
        resp = api_client.get(f'/api/jobs/{uuid.uuid4()}/gherkin/')
        assert resp.status_code == 404

    def test_gherkin_no_automated_uc_still_returns_zip(self, api_client, db):
        job = ConversionJob.objects.create(
            source_filename='no_auto.docx',
            word_file='uploads/word/no_auto.docx',
            status=ConversionJob.Status.DONE,
            use_cases_count=1,
        )
        ExtractedUseCase.objects.create(job=job, order=1, description='Solo test', is_automated=False)
        resp = api_client.get(f'/api/jobs/{job.id}/gherkin/')
        assert resp.status_code == 200
        assert resp.content[:2] == b'PK'


# ─────────────────────────────────────────────────────────────────────────────
# FILE SEARCH
# ─────────────────────────────────────────────────────────────────────────────
class TestFileSearchAPI:
    def test_short_query_returns_empty(self, api_client):
        resp = api_client.get('/api/files/search/?q=a')
        assert resp.status_code == 200
        assert resp.data['results'] == []

    def test_missing_query_returns_empty(self, api_client):
        resp = api_client.get('/api/files/search/')
        assert resp.status_code == 200
        assert resp.data['results'] == []

    def test_valid_query_returns_list(self, api_client):
        resp = api_client.get('/api/files/search/?q=test')
        assert resp.status_code == 200
        assert 'results' in resp.data
        assert isinstance(resp.data['results'], list)


# ─────────────────────────────────────────────────────────────────────────────
# JOB DELETE (Permission-based)
# ─────────────────────────────────────────────────────────────────────────────
@pytest.fixture
def job_owner(db):
    return User.objects.create_user('job_owner', 'jowner@test.com', 'pass1234')

@pytest.fixture
def team_owner_user(db):
    return User.objects.create_user('team_owner', 'town@test.com', 'pass1234')

@pytest.fixture
def team_admin_user(db):
    return User.objects.create_user('team_admin', 'tadmin@test.com', 'pass1234')

@pytest.fixture
def team_member_user(db):
    return User.objects.create_user('team_member', 'tmember@test.com', 'pass1234')

@pytest.fixture
def unaffiliated_user(db):
    return User.objects.create_user('unaffiliated', 'unaff@test.com', 'pass1234')

@pytest.fixture
def project_team(db, team_owner_user, team_admin_user, team_member_user, job_owner):
    t = Team.objects.create(name='Delete Test Team', owner=team_owner_user)
    TeamMember.objects.create(team=t, user=team_owner_user, role='owner')
    TeamMember.objects.create(team=t, user=team_admin_user, role='admin')
    TeamMember.objects.create(team=t, user=team_member_user, role='member')
    TeamMember.objects.create(team=t, user=job_owner, role='member')
    return t

@pytest.fixture
def team_project(db, project_team, team_owner_user):
    p = Project.objects.create(
        team=project_team, name='Delete Test Project', created_by=team_owner_user, visibility='public'
    )
    return p

@pytest.fixture
def personal_job(db, job_owner):
    return ConversionJob.objects.create(
        source_filename='personal.docx',
        word_file='uploads/word/personal.docx',
        status=ConversionJob.Status.DONE,
        uploaded_by=job_owner,
        use_cases_count=2,
    )

@pytest.fixture
def team_job(db, team_project, job_owner):
    return ConversionJob.objects.create(
        source_filename='team_recipe.docx',
        word_file='uploads/word/team_recipe.docx',
        status=ConversionJob.Status.DONE,
        project=team_project,
        uploaded_by=job_owner,
        use_cases_count=3,
    )


class TestJobDeleteAPI:
    """Tests for DELETE /api/jobs/<id>/ — permission-based recipe deletion."""

    def test_delete_personal_job_by_uploader(self, api_client, personal_job, job_owner):
        api_client.force_authenticate(user=job_owner)
        resp = api_client.delete(f'/api/jobs/{personal_job.id}/')
        assert resp.status_code == 204
        assert not ConversionJob.objects.filter(id=personal_job.id).exists()

    def test_delete_personal_job_by_other_user_returns_403(self, api_client, personal_job, unaffiliated_user):
        api_client.force_authenticate(user=unaffiliated_user)
        resp = api_client.delete(f'/api/jobs/{personal_job.id}/')
        assert resp.status_code == 403

    def test_delete_personal_job_unauthenticated_returns_403(self, api_client, personal_job):
        resp = api_client.delete(f'/api/jobs/{personal_job.id}/')
        assert resp.status_code == 403

    def test_delete_team_job_by_team_owner(self, api_client, team_job, team_owner_user):
        api_client.force_authenticate(user=team_owner_user)
        resp = api_client.delete(f'/api/jobs/{team_job.id}/')
        assert resp.status_code == 204
        assert not ConversionJob.objects.filter(id=team_job.id).exists()

    def test_delete_team_job_by_team_admin(self, api_client, team_job, team_admin_user):
        api_client.force_authenticate(user=team_admin_user)
        resp = api_client.delete(f'/api/jobs/{team_job.id}/')
        assert resp.status_code == 204
        assert not ConversionJob.objects.filter(id=team_job.id).exists()

    def test_delete_team_job_by_team_member_returns_403(self, api_client, team_job, team_member_user):
        api_client.force_authenticate(user=team_member_user)
        resp = api_client.delete(f'/api/jobs/{team_job.id}/')
        assert resp.status_code == 403

    def test_delete_team_job_by_unaffiliated_user_returns_403(self, api_client, team_job, unaffiliated_user):
        api_client.force_authenticate(user=unaffiliated_user)
        resp = api_client.delete(f'/api/jobs/{team_job.id}/')
        assert resp.status_code == 403

    def test_delete_team_job_unauthenticated_returns_403(self, api_client, team_job):
        resp = api_client.delete(f'/api/jobs/{team_job.id}/')
        assert resp.status_code == 403

    def test_delete_job_not_found(self, api_client, team_owner_user):
        api_client.force_authenticate(user=team_owner_user)
        resp = api_client.delete(f'/api/jobs/{uuid.uuid4()}/')
        assert resp.status_code == 404

    def test_cascade_deletes_use_cases(self, api_client, team_job, team_owner_user):
        uc = ExtractedUseCase.objects.create(job=team_job, order=1, description='Test UC')
        api_client.force_authenticate(user=team_owner_user)
        api_client.delete(f'/api/jobs/{team_job.id}/')
        assert not ExtractedUseCase.objects.filter(id=uc.id).exists()


# ─────────────────────────────────────────────────────────────────────────────
# UPLOAD
# ─────────────────────────────────────────────────────────────────────────────
@pytest.mark.django_db
class TestUploadAPI:
    def test_upload_requires_word_file(self, api_client):
        resp = api_client.post('/api/upload/', {}, format='multipart')
        assert resp.status_code == 400
        assert 'requis' in resp.data['error'].lower()

    def test_upload_rejects_invalid_docx(self, api_client):
        from django.core.files.uploadedfile import SimpleUploadedFile

        uploaded = SimpleUploadedFile('broken.docx', b'not a real docx')
        resp = api_client.post(
            '/api/upload/',
            {'word_file': uploaded},
            format='multipart',
        )
        assert resp.status_code == 400
        assert 'error' in resp.data
        assert ConversionJob.objects.count() == 0

    def test_upload_accepts_valid_docx(self, api_client, tmp_path):
        import io
        from django.core.files.uploadedfile import SimpleUploadedFile
        from docx import Document

        doc = Document()
        table = doc.add_table(rows=2, cols=3)
        table.rows[0].cells[0].text = 'Use Case'
        table.rows[0].cells[1].text = 'Description'
        table.rows[0].cells[2].text = 'Étapes'
        table.rows[1].cells[0].text = 'UC001'
        table.rows[1].cells[1].text = 'Connexion'
        table.rows[1].cells[2].text = 'Ouvrir la page'

        buffer = io.BytesIO()
        doc.save(buffer)
        uploaded = SimpleUploadedFile(
            'recette.docx',
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )

        resp = api_client.post(
            '/api/upload/',
            {'word_file': uploaded},
            format='multipart',
        )
        assert resp.status_code == 201
        assert resp.data['status'] == 'DONE'
        assert resp.data['use_cases_count'] == 1

    def test_upload_no_tables_returns_422(self, api_client):
        import io
        from django.core.files.uploadedfile import SimpleUploadedFile
        from docx import Document

        doc = Document()
        doc.add_paragraph('Document sans tableau')
        buffer = io.BytesIO()
        doc.save(buffer)
        uploaded = SimpleUploadedFile('empty.docx', buffer.getvalue())

        resp = api_client.post(
            '/api/upload/',
            {'word_file': uploaded},
            format='multipart',
        )
        assert resp.status_code == 422
        assert 'tableau' in resp.data['error'].lower()
        assert 'job_id' in resp.data
